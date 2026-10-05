"""Restaura la info legacy de Marine desde la referencia fija REF_LEGACY_202509.xlsx.

Las polizas de LIST_LEGACY no reciben BDX nuevo, asi que su Reserva Bruta y Deducible
deben ser identicos a la referencia (202509 MANUAL). Para cada archivo indicado:
  - filas legacy que existen en la referencia -> se copia GROSS RESERVE y DEDUCTIBLE de la referencia;
  - filas legacy que NO existen en la referencia (p. ej. pagos reasignados) -> se eliminan;
  - filas de la referencia que faltan en el archivo -> se insertan, tomando las demas
    columnas de la base plantilla (PROCESADO de 202510, previa a la corrupcion).
Las filas se emparejan por (poliza, claim, COVERAGE, pagado acumulado) porque los pagos
no cambian en legacy. Se respalda el archivo original antes de tocarlo.

Uso:
    python restaurar_legacy_desde_referencia.py            # simulacion (no escribe nada)
    python restaurar_legacy_desde_referencia.py --aplicar  # respalda y restaura
    python restaurar_legacy_desde_referencia.py --aplicar RUTA1 RUTA2 ...  # solo esos archivos
"""
import argparse
import datetime as dt
import os
import shutil
import sys

import pandas as pd

sys.path.insert(0, r'C:/Users/IKAL14/Documents/Integral/Marine/Diccionarios')
from config_marine import LIST_LEGACY, REF_LEGACY_PATH, _huella_legacy, _normalizar_claim

ONEDRIVE = r'C:/Users/IKAL14/OneDrive - Kot Insurance Company AG/Transporte, Carga y Embarcaciones'
MARINE = r'C:/Users/IKAL14/Documents/Integral/Marine'
PLANTILLA = f'{ONEDRIVE}/2025/202510/202510_Siniestros_Marine_PROCESADO.xlsx'
BACKUP_DIR = f'{MARINE}/_respaldo_legacy_{dt.date.today():%Y%m%d}'

# Desde que se corrompio (202512): cadena PROCESADO y salidas Final locales (xlsx).
# Los .pkl de Final no se tocan aqui (requieren pyarrow); se regeneran al correr el notebook 2.
ARCHIVOS = [
    f'{ONEDRIVE}/2025/202512/202512_Siniestros_Marine_PROCESADO.xlsx',
    f'{ONEDRIVE}/2026/202604/202604_Siniestros_Marine_PROCESADO.xlsx',
    f'{ONEDRIVE}/2026/202606/202606_Siniestros_Marine_PROCESADO.xlsx',
    f'{ONEDRIVE}/2026/202607/202607_Siniestros_Marine_PROCESADO.xlsx',
    f'{ONEDRIVE}/2026/202608/202608_Siniestros_Marine_PROCESADO.xlsx',
    f'{MARINE}/Procesados/202606/Final/202606_Siniestros_Marine.xlsx',
    f'{MARINE}/Procesados/202607/Final/202607_Siniestros_Marine.xlsx',
]

LEGACY = {str(p).strip() for p in LIST_LEGACY}


def _col_poliza(df):
    return [c for c in df.columns if str(c).startswith('INWARD POLICY N')][0]


def _con_llaves(df):
    """Agrega _POL, _KK (poliza|claim|coverage|pagado#ocurrencia) y marca las filas legacy."""
    df = df.copy()
    df['_POL'] = df[_col_poliza(df)].fillna('').astype(str).str.strip()
    paid = pd.to_numeric(df['Cumulative CLAIMS PAID'], errors='coerce').fillna(0).round(0).astype(int).astype(str)
    k = (df['_POL'] + '|' + df['CLAIM NUMBER'].map(_normalizar_claim) + '|'
         + df['COVERAGE'].fillna('').astype(str).str.strip().replace('NO ESPECIFICADO', '') + '|' + paid)
    df['_KK'] = k + '#' + k.groupby(k).cumcount().astype(str)
    df['_LEG'] = df['_POL'].isin(LEGACY)
    return df


def _leer(path):
    if path.lower().endswith('.pkl'):
        return pd.read_pickle(path)
    return pd.read_excel(path, dtype={'CLAIM NUMBER': str})


def planear(df, ref, plantilla):
    """Devuelve (actualizaciones {indice: (gross, ded)}, indices a borrar, filas a insertar)."""
    d = _con_llaves(df)
    leg = d[d['_LEG']]
    ref_idx = ref.set_index('_KK')
    en_ref = leg['_KK'].isin(ref_idx.index)
    borrar = list(leg.index[~en_ref])

    act = {}
    for i, r in leg[en_ref].iterrows():
        g, de = ref_idx.at[r['_KK'], 'GROSS RESERVE'], ref_idx.at[r['_KK'], 'DEDUCTIBLE']
        # tolerancia de 0.5 USD: ignora ruido de punto flotante de Excel (p. ej. 23910.63 vs .64)
        if abs(float(r['GROSS RESERVE']) - float(g)) > 0.5 or abs(float(r['DEDUCTIBLE']) - float(de)) > 0.5:
            act[i] = (g, de)

    faltan = ref_idx.index.difference(leg['_KK'])
    pl = _con_llaves(plantilla).set_index('_KK')
    filas = []
    for kk in faltan:
        if kk not in pl.index:
            raise ValueError(f'La fila {kk} falta en el archivo y tampoco esta en la plantilla; no se puede reconstruir.')
        fila = pl.loc[kk].copy()
        fila['GROSS RESERVE'] = ref_idx.at[kk, 'GROSS RESERVE']
        fila['DEDUCTIBLE'] = ref_idx.at[kk, 'DEDUCTIBLE']
        filas.append(fila)
    ins = pd.DataFrame(filas).reset_index(drop=True) if filas else pd.DataFrame()
    return act, borrar, ins


def _celda(v):
    if pd.isna(v):
        return None
    return v.to_pydatetime() if isinstance(v, pd.Timestamp) else v


def aplicar_xlsx(path, df, act, borrar, ins):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    cols = {c.value: c.column for c in ws[1] if c.value is not None}
    if ws.max_row - 1 != len(df):
        raise ValueError(f'{path}: la hoja tiene {ws.max_row - 1} filas y pandas leyo {len(df)}; no se edita.')
    for i, (g, de) in act.items():
        ws.cell(row=i + 2, column=cols['GROSS RESERVE']).value = float(g)
        ws.cell(row=i + 2, column=cols['DEDUCTIBLE']).value = float(de)
    for i in sorted(borrar, reverse=True):
        ws.delete_rows(i + 2)
    for _, f in ins.iterrows():
        ws.append([_celda(f[c]) if c in f.index else None for c in cols])
    wb.save(path)


def aplicar_pkl(path, df, act, borrar, ins):
    out = df.copy()
    for i, (g, de) in act.items():
        out.loc[i, 'GROSS RESERVE'], out.loc[i, 'DEDUCTIBLE'] = g, de
    out = out.drop(index=borrar)
    if len(ins):
        out = pd.concat([out, ins.reindex(columns=out.columns)], ignore_index=True)
    out.reset_index(drop=True).to_pickle(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--aplicar', action='store_true', help='respalda y escribe; sin esto solo simula')
    ap.add_argument('archivos', nargs='*')
    a = ap.parse_args()

    ref = _con_llaves(pd.read_excel(REF_LEGACY_PATH, dtype={'CLAIM NUMBER': str}))
    ref = ref[ref['_LEG']]
    assert ref['_KK'].is_unique
    plantilla = _leer(PLANTILLA)
    h_ref = _huella_legacy(ref.rename(columns={_col_poliza(ref): 'INWARD POLICY N°'}))

    for path in (a.archivos or ARCHIVOS):
        print(f'\n== {path}')
        if not os.path.exists(path):
            print('   (no existe, se omite)')
            continue
        try:
            df = _leer(path)
        except PermissionError:
            print('   ERROR: archivo abierto/bloqueado (cierralo en Excel y vuelve a correr). Se omite.')
            continue
        df = df.rename(columns={_col_poliza(df): 'INWARD POLICY N°'})
        act, borrar, ins = planear(df, ref, plantilla)
        pagado_antes = pd.to_numeric(df.loc[_con_llaves(df)['_LEG'], 'Cumulative CLAIMS PAID'], errors='coerce').sum()
        pols = lambda idx: sorted(_con_llaves(df).loc[idx, '_POL'].unique()) if len(idx) else []
        print(f'   actualizar Reserva/Deducible: {len(act)} fila(s) {pols(list(act))}')
        print(f'   eliminar (no estan en la referencia): {len(borrar)} fila(s) {pols(borrar)}')
        print(f'   insertar (faltan vs referencia): {len(ins)} fila(s)')
        if not (act or borrar or len(ins)):
            continue
        if not a.aplicar:
            continue
        os.makedirs(BACKUP_DIR, exist_ok=True)
        resp = os.path.join(BACKUP_DIR, os.path.basename(os.path.dirname(path)) + '__' + os.path.basename(path))
        shutil.copy2(path, resp)
        try:
            (aplicar_pkl if path.lower().endswith('.pkl') else aplicar_xlsx)(path, df, act, borrar, ins)
        except PermissionError:
            print('   ERROR: archivo bloqueado al guardar. Se omite (el respaldo ya se copio).')
            continue
        nuevo = _leer(path)
        nuevo = nuevo.rename(columns={_col_poliza(nuevo): 'INWARD POLICY N°'})
        ok_ref = (_huella_legacy(nuevo) == h_ref)
        pagado_despues = pd.to_numeric(nuevo.loc[_con_llaves(nuevo)['_LEG'], 'Cumulative CLAIMS PAID'], errors='coerce').sum()
        print(f'   respaldo: {resp}')
        print(f'   verificacion vs referencia: {"OK" if ok_ref else "FALLA"} | '
              f'pagado legacy antes/despues: {pagado_antes:,.2f} / {pagado_despues:,.2f}')


if __name__ == '__main__':
    main()
